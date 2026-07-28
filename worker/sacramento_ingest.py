"""Sacramento-only OCR and Qdrant Cloud Inference ingestion worker."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import tempfile
import uuid
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from pypdf import PdfReader
from qdrant_client import QdrantClient, models


DENSE_MODEL = "sentence-transformers/all-minilm-l6-v2"
SPARSE_MODEL = "qdrant/bm25"
NAMESPACE = uuid.UUID("9dd5ce43-41e3-5de0-9aa0-d635cb41c9ee")
COLLECTIONS = {
    "public_foia": ("lpt_research_v2_384", "lpt_research_live"),
    "prior_appeal": ("appeal_comps_v2_384", "appeal_comps_live"),
    "private_case": ("case_private_v1_384", "case_private_live"),
}


@dataclass(frozen=True)
class PageText:
    page: int
    text: str


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--bootstrap", action="store_true")
    parser.add_argument("--activate-aliases", action="store_true")
    parser.add_argument(
        "--visibility",
        choices=tuple(COLLECTIONS),
        default="public_foia",
    )
    parser.add_argument("--appeal-id")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--process-private-queue", action="store_true")
    parser.add_argument("--process-appeal-comps", action="store_true")
    return parser.parse_args()


def client() -> QdrantClient:
    url = os.environ.get("QDRANT_URL", "").rstrip("/")
    api_key = os.environ.get("QDRANT_API_KEY", "")
    if not url or not api_key:
        raise RuntimeError("QDRANT_URL and QDRANT_API_KEY are required.")
    return QdrantClient(
        url=url,
        api_key=api_key,
        cloud_inference=True,
        timeout=60,
    )


def bootstrap(qdrant: QdrantClient, activate_aliases: bool) -> None:
    existing = {item.name for item in qdrant.get_collections().collections}
    for collection, _alias in COLLECTIONS.values():
        if collection not in existing:
            qdrant.create_collection(
                collection_name=collection,
                vectors_config={
                    "dense": models.VectorParams(
                        size=384,
                        distance=models.Distance.COSINE,
                    )
                },
                sparse_vectors_config={
                    "bm25": models.SparseVectorParams(
                        modifier=models.Modifier.IDF,
                    )
                },
            )
        payload_schema = qdrant.get_collection(collection).payload_schema or {}
        for field, schema in (
            ("county_slug", models.PayloadSchemaType.KEYWORD),
            ("visibility", models.PayloadSchemaType.KEYWORD),
            ("appeal_id", models.PayloadSchemaType.KEYWORD),
            ("tax_year", models.PayloadSchemaType.KEYWORD),
            ("property_type", models.PayloadSchemaType.KEYWORD),
            ("document_type", models.PayloadSchemaType.KEYWORD),
            ("outcome", models.PayloadSchemaType.KEYWORD),
        ):
            if field in payload_schema:
                continue
            qdrant.create_payload_index(
                collection_name=collection,
                field_name=field,
                field_schema=schema,
                wait=True,
            )
    if activate_aliases:
        switch_aliases(qdrant)


def switch_aliases(qdrant: QdrantClient) -> None:
    current = {
        item.alias_name: item.collection_name
        for item in qdrant.get_aliases().aliases
    }
    operations: list[models.AliasOperations] = []
    for collection, alias in COLLECTIONS.values():
        if current.get(alias) == collection:
            continue
        if alias in current:
            operations.append(
                models.DeleteAliasOperation(
                    delete_alias=models.DeleteAlias(alias_name=alias)
                )
            )
        operations.append(
            models.CreateAliasOperation(
                create_alias=models.CreateAlias(
                    collection_name=collection,
                    alias_name=alias,
                )
            )
        )
    if operations:
        qdrant.update_collection_aliases(change_aliases_operations=operations)


def ingest_manifest(
    qdrant: QdrantClient,
    path: Path,
    visibility: str,
    appeal_id: str | None,
    dry_run: bool,
) -> dict[str, int]:
    if visibility == "private_case" and not appeal_id:
        raise RuntimeError("--appeal-id is required for private-case ingestion.")
    raw = json.loads(path.read_text(encoding="utf-8"))
    documents = raw if isinstance(raw, list) else raw.get("documents", [])
    ocr_candidates = [] if isinstance(raw, list) else raw.get("skipped", [])
    totals = {"documents": 0, "pages": 0, "points": 0, "ocr_documents": 0}
    collection = COLLECTIONS[visibility][0]

    for document in documents:
        pages = pages_from_document(document)
        points = points_for_document(document, pages, visibility, appeal_id)
        totals["documents"] += 1
        totals["pages"] += len(pages)
        totals["points"] += len(points)
        if not dry_run:
            upsert_batches(qdrant, collection, points)
            record_ingestion_job(
                document,
                collection,
                pages,
                points,
                "indexed",
                visibility=visibility,
            )

    for candidate in ocr_candidates:
        local_path = candidate.get("localPath")
        if not local_path:
            continue
        document = {
            "sourceId": candidate.get("id"),
            "title": candidate.get("name"),
            "sourceUrl": candidate.get("sourceUrl"),
            "modifiedAt": candidate.get("modifiedAt"),
            "metadata": {"ocr": True},
        }
        try:
            pages = extract_pdf_pages(Path(local_path), force_ocr=True)
            points = points_for_document(document, pages, visibility, appeal_id)
            totals["documents"] += 1
            totals["ocr_documents"] += 1
            totals["pages"] += len(pages)
            totals["points"] += len(points)
            if not dry_run:
                upsert_batches(qdrant, collection, points)
                record_ingestion_job(
                    document,
                    collection,
                    pages,
                    points,
                    "indexed",
                    visibility=visibility,
                )
        except Exception as error:  # noqa: BLE001 - continue remaining OCR files
            if not dry_run:
                record_ingestion_job(
                    document,
                    collection,
                    [],
                    [],
                    "failed",
                    "ocr_failed",
                    str(error),
                    visibility=visibility,
                )
    return totals


def process_private_queue(
    qdrant: QdrantClient,
    limit: int = 50,
) -> dict[str, int]:
    rows = supabase_json(
        "GET",
        (
            "rest/v1/case_document_vectors"
            "?status=in.(queued,failed)"
            "&select=id,appeal_id,document_id,content_hash"
            f"&limit={limit}"
        ),
    )
    totals = {"seen": len(rows), "indexed": 0, "failed": 0}
    for mapping in rows:
        mapping_id = mapping["id"]
        try:
            supabase_json(
                "PATCH",
                f"rest/v1/case_document_vectors?id=eq.{mapping_id}",
                {"status": "extracting", "error_message": None},
            )
            documents = supabase_json(
                "GET",
                (
                    "rest/v1/appeal_documents"
                    f"?id=eq.{mapping['document_id']}"
                    "&select=id,title,doc_type,file_path"
                    "&limit=1"
                ),
            )
            if not documents:
                raise RuntimeError("Private document metadata was not found.")
            document = documents[0]
            suffix = Path(document["file_path"]).suffix or ".pdf"
            with tempfile.TemporaryDirectory(prefix="lpt-private-") as temp:
                local_path = Path(temp) / f"source{suffix}"
                local_path.write_bytes(
                    supabase_bytes(
                        "GET",
                        "storage/v1/object/case-documents/"
                        + urllib.parse.quote(document["file_path"], safe="/"),
                    )
                )
                pages = extract_private_pages(local_path)
            source = {
                "sourceId": document["id"],
                "title": document["title"],
                "documentType": document["doc_type"],
                "metadata": {
                    "content_hash": mapping["content_hash"],
                },
            }
            points = points_for_document(
                source,
                pages,
                "private_case",
                mapping["appeal_id"],
            )
            upsert_batches(
                qdrant,
                COLLECTIONS["private_case"][0],
                points,
            )
            supabase_json(
                "PATCH",
                f"rest/v1/case_document_vectors?id=eq.{mapping_id}",
                {
                    "status": "indexed",
                    "qdrant_collection": COLLECTIONS["private_case"][0],
                    "point_ids": [str(point.id) for point in points],
                    "indexed_at": utc_now(),
                    "error_message": None,
                },
            )
            totals["indexed"] += 1
        except Exception as error:  # noqa: BLE001 - each queue item must continue
            totals["failed"] += 1
            supabase_json(
                "PATCH",
                f"rest/v1/case_document_vectors?id=eq.{mapping_id}",
                {
                    "status": "failed",
                    "error_message": str(error)[:2000],
                },
            )
    return totals


def process_appeal_comps(qdrant: QdrantClient) -> dict[str, int]:
    totals = {"seen": 0, "indexed": 0, "skipped": 0}
    offset = 0
    while True:
        appeals = supabase_json(
            "GET",
            (
                "rest/v1/appeals"
                "?source_system=eq.sacramento_lambda"
                "&select=id,appeal_number,tax_year,status,source_status,enrolled_value,"
                "claimed_value,final_value,reduction_pct,outcome_notes,property_id"
                f"&limit=500&offset={offset}"
            ),
        )
        if not appeals:
            break
        property_ids = sorted(
            {str(item["property_id"]) for item in appeals if item.get("property_id")}
        )
        properties: list[dict[str, Any]] = []
        for start in range(0, len(property_ids), 100):
            values = ",".join(property_ids[start : start + 100])
            properties.extend(
                supabase_json(
                    "GET",
                    (
                        "rest/v1/properties"
                        f"?id=in.({values})"
                        "&select=id,county,property_type,address,city,apn"
                    ),
                )
            )
        property_by_id = {str(item["id"]): item for item in properties}
        points: list[models.PointStruct] = []
        for appeal in appeals:
            totals["seen"] += 1
            property_record = property_by_id.get(str(appeal.get("property_id")))
            county = str((property_record or {}).get("county") or "")
            if county.lower().replace(" county", "").strip() != "sacramento":
                totals["skipped"] += 1
                continue
            outcome = appeal_outcome(appeal)
            text = " ".join(
                str(value)
                for value in (
                    f"Sacramento appeal {appeal.get('appeal_number') or appeal['id']}.",
                    f"Tax year {appeal.get('tax_year')}.",
                    f"Property type {(property_record or {}).get('property_type') or 'unknown'}.",
                    f"Status or outcome {outcome}.",
                    appeal.get("outcome_notes") or "",
                )
                if value
            )
            document = {
                "sourceId": str(appeal["id"]),
                "title": f"Appeal {appeal.get('appeal_number') or appeal['id']}",
                "documentType": "appeal_record",
                "metadata": {
                    "appeal_number": appeal.get("appeal_number"),
                    "tax_year": scalar_text(appeal.get("tax_year")),
                    "property_type": (property_record or {}).get("property_type"),
                    "outcome": scalar_text(outcome),
                    "reason": appeal.get("outcome_notes"),
                    "reduction_pct": appeal.get("reduction_pct"),
                    "enrolled_value": appeal.get("enrolled_value"),
                    "claimed_value": appeal.get("claimed_value"),
                    "final_value": appeal.get("final_value"),
                },
            }
            points.extend(
                points_for_document(
                    document,
                    [PageText(page=1, text=text)],
                    "prior_appeal",
                    None,
                )
            )
        upsert_batches(qdrant, COLLECTIONS["prior_appeal"][0], points)
        totals["indexed"] += len(points)
        if len(appeals) < 500:
            break
        offset += 500
    return totals


def pages_from_document(document: dict[str, Any]) -> list[PageText]:
    local_path = document.get("localPath") or document.get("file_path")
    if local_path and Path(local_path).suffix.lower() == ".pdf":
        return extract_pdf_pages(Path(local_path))
    text = str(document.get("text") or "").strip()
    if not text:
        return []
    return [PageText(page=int(document.get("page", 1)), text=text)]


def extract_pdf_pages(path: Path, force_ocr: bool = False) -> list[PageText]:
    pages = read_pdf(path)
    character_count = sum(len(item.text) for item in pages)
    if not force_ocr and character_count >= max(500, len(pages) * 80):
        return pages
    with tempfile.TemporaryDirectory(prefix="lpt-ocr-") as temp:
        output = Path(temp) / "ocr.pdf"
        command = [
            "ocrmypdf",
            "--skip-text",
            "--deskew",
            "--rotate-pages",
            "--output-type",
            "pdf",
            str(path),
            str(output),
        ]
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            timeout=900,
        )
        if completed.returncode not in (0, 6):
            raise RuntimeError(
                f"OCR failed for {path.name}: {completed.stderr[-500:]}"
            )
        return read_pdf(output)


def extract_private_pages(path: Path) -> list[PageText]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_pages(path)
    if suffix in (".txt", ".csv"):
        return [
            PageText(
                page=1,
                text=path.read_text(encoding="utf-8", errors="replace"),
            )
        ]
    if suffix == ".docx":
        from docx import Document as WordDocument

        document = WordDocument(str(path))
        text = "\n".join(
            paragraph.text for paragraph in document.paragraphs if paragraph.text
        )
        return [PageText(page=1, text=text)]
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        pages: list[PageText] = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows = [
                "\t".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            pages.append(
                PageText(
                    page=index,
                    text=f"Worksheet: {sheet.title}\n" + "\n".join(rows),
                )
            )
        return pages
    if suffix in (".png", ".jpg", ".jpeg"):
        completed = subprocess.run(
            ["tesseract", str(path), "stdout", "-l", "eng"],
            check=False,
            capture_output=True,
            text=True,
            timeout=300,
        )
        if completed.returncode != 0:
            raise RuntimeError(f"Image OCR failed: {completed.stderr[-500:]}")
        return [PageText(page=1, text=completed.stdout)]
    raise RuntimeError(f"Private worker does not support {suffix} extraction.")


def read_pdf(path: Path) -> list[PageText]:
    reader = PdfReader(str(path))
    pages: list[PageText] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(PageText(page=index, text=text))
    return pages


def points_for_document(
    document: dict[str, Any],
    pages: Iterable[PageText],
    visibility: str,
    appeal_id: str | None,
) -> list[models.PointStruct]:
    source_id = str(document.get("sourceId") or document.get("document_id") or "")
    title = str(document.get("title") or "Sacramento County record")
    if not source_id:
        raise RuntimeError(f"Document '{title}' has no stable source ID.")
    metadata = dict(document.get("metadata") or {})
    points: list[models.PointStruct] = []
    for page in pages:
        text = page.text
        if visibility == "private_case":
            text = redact(text)
        for chunk_index, chunk in enumerate(chunk_text(text)):
            content_hash = hashlib.sha256(chunk.encode("utf-8")).hexdigest()
            point_id = str(
                uuid.uuid5(
                    NAMESPACE,
                    f"{source_id}:{page.page}:{chunk_index}:{content_hash}",
                )
            )
            payload = {
                **metadata,
                "text": chunk,
                "title": title,
                "county": "Sacramento",
                "county_slug": "sacramento",
                "document_id": source_id,
                "document_type": document.get("documentType", "other"),
                "page_start": page.page,
                "page_end": page.page,
                "content_hash": content_hash,
                "source_url": (
                    None if visibility == "private_case" else document.get("sourceUrl")
                ),
                "source_modified_at": document.get("modifiedAt"),
                "tax_year": scalar_text(metadata.get("tax_year")),
                "property_type": scalar_text(metadata.get("property_type")),
                "outcome": scalar_text(metadata.get("outcome")),
                "visibility": visibility,
                "source_type": (
                    "public" if visibility == "public_foia" else visibility
                ),
                "appeal_id": appeal_id,
                "redaction_version": (
                    "pii-v1" if visibility == "private_case" else None
                ),
            }
            points.append(
                models.PointStruct(
                    id=point_id,
                    payload=payload,
                    vector={
                        "dense": models.Document(text=chunk, model=DENSE_MODEL),
                        "bm25": models.Document(text=chunk, model=SPARSE_MODEL),
                    },
                )
            )
    return points


def chunk_text(text: str, size: int = 1500, overlap: int = 180) -> list[str]:
    clean = re.sub(r"[ \t]+", " ", text).strip()
    if not clean:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + size)
        if end < len(clean):
            split = clean.rfind(" ", start + int(size * 0.6), end)
            if split > start:
                end = split
        chunks.append(clean[start:end].strip())
        if end >= len(clean):
            break
        start = max(start + 1, end - overlap)
    return chunks


def redact(text: str) -> str:
    rules = (
        (r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", "[REDACTED EMAIL]"),
        (
            r"(?<!\d)(?:\+?1[\s.-]?)?(?:\(\d{3}\)|\d{3})[\s.-]\d{3}[\s.-]\d{4}(?!\d)",
            "[REDACTED PHONE]",
        ),
        (r"(?<!\d)\d{3}-\d{2}-\d{4}(?!\d)", "[REDACTED SSN]"),
        (
            r"\b(owner|client|taxpayer|contact|applicant)\s*(name)?\s*:\s*[^\n,;]{2,100}",
            r"\1: [REDACTED NAME]",
        ),
    )
    result = text
    for pattern, replacement in rules:
        result = re.sub(pattern, replacement, result, flags=re.IGNORECASE)
    return result


def scalar_text(value: Any) -> str | None:
    if value is None:
        return None
    clean = str(value).strip()
    return clean or None


def appeal_outcome(appeal: dict[str, Any]) -> str:
    status = str(
        appeal.get("source_status") or appeal.get("status") or ""
    ).lower()
    if "withdraw" in status:
        return "Withdrawn"
    enrolled = numeric(appeal.get("enrolled_value"))
    final = numeric(appeal.get("final_value"))
    if enrolled is not None and final is not None:
        return "Win" if final < enrolled else "Loss"
    if any(term in status for term in ("approved", "reduced", "win", "granted")):
        return "Win"
    if any(term in status for term in ("denied", "loss", "upheld")):
        return "Loss"
    return "Pending"


def numeric(value: Any) -> float | None:
    try:
        return None if value is None else float(value)
    except (TypeError, ValueError):
        return None


def upsert_batches(
    qdrant: QdrantClient,
    collection: str,
    points: list[models.PointStruct],
) -> None:
    for start in range(0, len(points), 64):
        qdrant.upsert(
            collection_name=collection,
            points=points[start : start + 64],
            wait=True,
        )


def record_ingestion_job(
    document: dict[str, Any],
    collection: str,
    pages: Iterable[PageText],
    points: list[models.PointStruct],
    status: str,
    error_code: str | None = None,
    error_message: str | None = None,
    visibility: str = "public_foia",
) -> None:
    if not os.environ.get("SUPABASE_URL") or not os.environ.get(
        "SUPABASE_SECRET_KEY"
    ):
        return
    source_id = str(document.get("sourceId") or document.get("document_id") or "")
    metadata = dict(document.get("metadata") or {})
    page_list = list(pages)
    content_hash = metadata.get("content_hash") or hashlib.sha256(
        "\n".join(page.text for page in page_list).encode("utf-8")
    ).hexdigest()
    sources = supabase_json(
        "POST",
        (
            "rest/v1/foia_ingestion_sources"
            "?on_conflict=source_type,external_id"
        ),
        {
            "source_type": (
                "google_drive" if visibility == "public_foia" else "manual"
            ),
            "external_id": source_id,
            "title": document.get("title") or "Sacramento County record",
            "county": "Sacramento",
            "county_slug": "sacramento",
            "document_type": document.get("documentType") or "other",
            "source_url": document.get("sourceUrl"),
            "visibility": visibility,
            "content_hash": content_hash,
            "source_modified_at": document.get("modifiedAt"),
            "metadata": metadata,
        },
    )
    if not sources:
        raise RuntimeError("Supabase did not return the ingestion source.")
    supabase_json(
        "POST",
        "rest/v1/foia_ingestion_jobs",
        {
            "source_id": sources[0]["id"],
            "status": status,
            "attempt_count": 1,
            "qdrant_collection": collection,
            "page_count": len(page_list),
            "chunk_count": len(points),
            "qdrant_point_ids": [str(point.id) for point in points],
            "error_code": error_code,
            "error_message": None if error_message is None else error_message[:2000],
            "started_at": utc_now(),
            "finished_at": utc_now(),
        },
    )


def supabase_json(
    method: str,
    path: str,
    body: Any | None = None,
) -> Any:
    data = None if body is None else json.dumps(body).encode("utf-8")
    raw = supabase_request(method, path, data)
    return json.loads(raw.decode("utf-8")) if raw else []


def supabase_bytes(method: str, path: str) -> bytes:
    return supabase_request(method, path, None)


def supabase_request(method: str, path: str, data: bytes | None) -> bytes:
    base_url = os.environ.get("SUPABASE_URL", "").rstrip("/")
    secret = os.environ.get("SUPABASE_SECRET_KEY", "")
    if not base_url or not secret:
        raise RuntimeError("SUPABASE_URL and SUPABASE_SECRET_KEY are required.")
    request = urllib.request.Request(
        f"{base_url}/{path}",
        method=method,
        data=data,
        headers={
            "apikey": secret,
            "Authorization": f"Bearer {secret}",
            "Content-Type": "application/json",
            "Prefer": "resolution=merge-duplicates,return=representation",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            return response.read()
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(
            f"Supabase request failed ({error.code}): {detail[:500]}"
        ) from error


def utc_now() -> str:
    from datetime import datetime, timezone

    return datetime.now(timezone.utc).isoformat()


def main() -> None:
    args = arguments()
    qdrant = client()
    if args.bootstrap:
        bootstrap(qdrant, args.activate_aliases)
    elif args.activate_aliases:
        switch_aliases(qdrant)
    if args.manifest:
        totals = ingest_manifest(
            qdrant,
            args.manifest.resolve(),
            args.visibility,
            args.appeal_id,
            args.dry_run,
        )
        print(json.dumps(totals, indent=2))
    if args.process_private_queue:
        print(json.dumps(process_private_queue(qdrant), indent=2))
    if args.process_appeal_comps:
        print(json.dumps(process_appeal_comps(qdrant), indent=2))


if __name__ == "__main__":
    main()
